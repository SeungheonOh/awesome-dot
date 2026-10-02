#!/usr/bin/env python3
"""Create and check this bounded fictional XLSX example. Never overwrite a run.

Requires existing Python 3.10+, openpyxl, XlsxWriter, LibreOffice and pdftoppm.
No installs, network access, macros or external data refreshes are used.
"""

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import zipfile
from collections import Counter
from copy import copy
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import xlsxwriter


HERE = Path(__file__).resolve().parent
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def require(ok, message):
    if not ok:
        raise RuntimeError(message)


def digest(path):
    return {"bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def dump(path, value):
    with path.open("x", encoding="utf-8") as stream:
        json.dump(value, stream, indent=2, ensure_ascii=False)
        stream.write("\n")


def run_command(args, operations):
    completed = subprocess.run(args, text=True, capture_output=True, timeout=120)
    operation = {"command": list(map(str, args)), "returncode": completed.returncode,
                 "stdout": completed.stdout.strip(), "stderr": completed.stderr.strip()}
    operations.append(operation)
    require(completed.returncode == 0, "Command failed: " + json.dumps(operation))
    return completed.stdout.strip()


def consumer_convert(input_path, output_dir, extension, profile, operations, soffice):
    output_dir.mkdir()
    profile.mkdir()
    # Ordinary conversion can preserve stale formula caches. This isolated profile
    # asks Calc to recalculate OOXML formulas on load (RECALC_ALWAYS = 0).
    # It does not change the operator's LibreOffice preferences.
    user_profile = profile / "user"
    user_profile.mkdir()
    (user_profile / "registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
        '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>'
        '</item></oor:items>\n', encoding="utf-8")
    result = output_dir / (input_path.stem + "." + extension)
    require(not result.exists(), f"Refusing existing output: {result}")
    format_arg = "xlsx:Calc MS Excel 2007 XML" if extension == "xlsx" else "pdf:calc_pdf_Export"
    run_command([soffice, "-env:UserInstallation=" + profile.as_uri(),
                 "--headless", "--convert-to", format_arg,
                 "--outdir", str(output_dir), str(input_path)], operations)
    require(result.is_file() and result.stat().st_size > 0,
            f"Consumer reported success but did not create {result}")
    return result


def make_source(path, fixture):
    require(not path.exists(), f"Refusing existing output: {path}")
    book = xlsxwriter.Workbook(path, {"strings_to_formulas": False,
                                    "strings_to_urls": False})
    book.set_properties({"title": "Fictional equipment export",
                         "subject": "Inventory cleanup and formula preservation",
                         "author": "awesome-dot fictional example"})
    book.set_calc_mode("auto")
    overview = book.add_worksheet("Overview")
    inventory = book.add_worksheet("Inventory")
    base = {"font_name": "Liberation Sans", "font_size": 11,
            "font_color": "#172A3A", "valign": "vcenter"}
    body = book.add_format(base)
    title = book.add_format({**base, "font_size": 16, "bold": True})
    note = book.add_format({**base, "font_size": 10, "font_color": "#526273"})
    header = book.add_format({**base, "font_color": "white", "bg_color": "#254B65",
                              "bold": True, "text_wrap": True, "align": "center"})
    identifier = book.add_format({**base, "num_format": "@"})
    integer = book.add_format({**base, "num_format": "0", "align": "right", "indent": 1})
    calculation = book.add_format({**base, "num_format": "0", "align": "right", "indent": 1,
                                   "bg_color": "#EDF3F6"})
    status = book.add_format({**base, "bold": True})
    warning = book.add_format({"bg_color": "#FFF1CF", "font_color": "#6B4D08"})
    shortage = book.add_format({"bg_color": "#FCE1DF", "font_color": "#862D2A"})
    for sheet in (overview, inventory):
        sheet.hide_gridlines(2)
        sheet.set_default_row(24)
        sheet.set_landscape()
        sheet.set_paper(9)
        sheet.fit_to_pages(1, 1)
        sheet.set_margins(0.35, 0.35, 0.40, 0.40)
        sheet.set_zoom(90)
    overview.set_tab_color("#254B65")
    overview.set_column("A:A", 43, body)
    overview.set_column("B:B", 18, body)
    overview.set_column("C:C", 54, note)
    overview.write("A2", "Equipment availability review", title)
    overview.set_row(1, 28)
    overview.write("A3", "Fictional export. Counts are equipment units.", note)
    metrics = [
        ("Numeric row-availability subtotal", "=SUM(AvailableUnits)", "Includes both disputed 0044 rows"),
        ("Rows without numeric counts", '=COUNTIF(Inventory!F8:F13,"unknown")', "Missing counts remain visible"),
        ("Export records retained", "=COUNTA(Inventory!A8:A13)", "No records removed or combined"),
        ("Rows for kit 0044", '=COUNTIF(Inventory!B8:B13,"0044")', "Two records have different ready counts"),
        ("Numeric subtotal from kit 0044", '=SUMIF(Inventory!B8:B13,"0044",Inventory!F8:F13)', "Included in the row subtotal; unresolved"),
        ("Numeric subtotal outside kit 0044", "=B4-B8", "Row availability outside the disputed pair"),
        ("Numeric ready-count cells", "=COUNT(Inventory!D8:D13)", "Counts numeric zero; excludes blanks and text"),
    ]
    for row, (label, formula, explanation) in enumerate(metrics, 3):
        overview.write(row, 0, label, body)
        overview.write_formula(row, 1, formula, calculation)
        overview.write(row, 2, explanation, note)
    overview.write("A12", "Counts completeness", body)
    overview.write_formula("B12", '=IF(B5>0,"Incomplete","Complete")', status)
    overview.write("C12", "Filling a count does not settle the key conflict", note)
    overview.write("A13", "Kit 0044 resolution", body)
    overview.write_formula("B13", '=IF(B7>1,"Unresolved","No repeat")', status)
    overview.write("C13", "No rule identifies a survivor", note)
    overview.write("A16", "Evidence still needed", status)
    overview.write("A17", "0025 and 0031: obtain the missing ready counts.", body)
    overview.write("A18", "0044: establish whether both rows are valid or which record is authoritative.", body)
    overview.write("A20", "Actual usable inventory remains unresolved: missing counts and conflicting rows.", note)
    overview.print_area("A1:C21")

    widths = [16, 12, 19, 11, 12, 12, 16, 37]
    for col, width in enumerate(widths):
        inventory.set_column(col, col, width, body)
    inventory.write("A2", "Workshop equipment export", title)
    inventory.set_row(1, 28)
    inventory.write("A3", "Fictional fixture. One row is one export record; kit ID is not a unique row key.", note)
    inventory.write("A4", "Blank and N/A mean unknown counts. Numeric 0 means none.", note)
    inventory.write("A5", "Available units = ready - reserved when both inputs are numeric.", note)
    headers = ["Source row", "Kit ID", "Equipment", "Ready units", "Reserved units",
               "Available units", "Status", "Review reason"]
    inventory.write_row("A7", headers, header)
    inventory.set_row(6, 34)
    for ordinal, values in enumerate(fixture["source_rows"], 8):
        source, kit, item, ready, reserved, state = values
        inventory.write_string(ordinal - 1, 0, source, identifier)
        inventory.write_string(ordinal - 1, 1, kit, identifier)
        inventory.write_string(ordinal - 1, 2, item, body)
        for col, value in ((3, ready), (4, reserved)):
            if value is None:
                inventory.write_blank(ordinal - 1, col, None, integer)
            elif isinstance(value, str):
                inventory.write_string(ordinal - 1, col, value, integer)
            else:
                inventory.write_number(ordinal - 1, col, value, integer)
        formula = (f'=IF(OR(NOT(ISNUMBER(D{ordinal})),NOT(ISNUMBER(E{ordinal}))),'
                   f'"unknown",D{ordinal}-E{ordinal})')
        inventory.write_formula(ordinal - 1, 5, formula, calculation)
        inventory.write_string(ordinal - 1, 6, state, identifier)
        inventory.write_blank(ordinal - 1, 7, None, body)
    inventory.write("A16", "Keep both 0044 rows until the source owner supplies a record decision.", note)
    inventory.freeze_panes(7, 2)
    inventory.autofilter("A7:H13")
    inventory.data_validation("D8:E13", {"validate": "integer", "criteria": ">=",
                                         "value": 0, "ignore_blank": True,
                                         "error_title": "Use a whole-unit count",
                                         "error_message": "Enter zero or a positive whole number, or leave unknown blank."})
    inventory.conditional_format("F8:F13", {"type": "cell", "criteria": "==",
                                             "value": '"unknown"', "format": warning})
    inventory.conditional_format("F8:F13", {"type": "cell", "criteria": "<",
                                             "value": 0, "format": shortage})
    inventory.conditional_format("H8:H13", {"type": "formula", "criteria": '=$H8<>""',
                                             "format": warning})
    inventory.print_area("A1:H17")
    book.define_name("AvailableUnits", "=Inventory!$F$8:$F$13")
    book.close()


def transform(source, target, fixture):
    require(not target.exists(), f"Refusing existing output: {target}")
    book = openpyxl.load_workbook(source, data_only=False)
    sheet = book["Inventory"]
    ledger = []

    def change(address, value, rule):
        cell = sheet[address]
        if cell.value != value or type(cell.value) is not type(value):
            ledger.append({"cell": address, "before": cell.value,
                           "after": value, "rule": rule})
            cell.value = value

    ids = Counter(sheet.cell(row, 2).value.strip() for row in range(8, 14))
    for row in range(8, 14):
        kit = sheet.cell(row, 2).value.strip()
        change(f"B{row}", kit, "ID_PADDING")
        ready = sheet.cell(row, 4).value
        reason = None
        if ready is None:
            reason = "Missing ready count: source blank"
        elif ready == "N/A":
            change(f"D{row}", None, "MISSING_MARKER")
            reason = "Missing count: source_marker (N/A)"
        elif isinstance(ready, str):
            require(ready.strip().isascii() and ready.strip().isdigit(),
                    f"Unrecognized count at D{row}; do not guess")
            change(f"D{row}", int(ready.strip()), "WHOLE_UNIT_TEXT")
        state = sheet.cell(row, 7).value.strip().casefold()
        require(state in fixture["rules"]["status"], f"Unmapped status at G{row}")
        change(f"G{row}", fixture["rules"]["status"][state], "STATUS_MAP")
        if ids[kit] > 1:
            require(reason is None, "This bounded fixture has separate missing and conflicting rows")
            reason = f"Conflicting kit ID {kit}"
        if reason:
            change(f"H{row}", reason, "REVIEW_EVIDENCE")
    book.save(target)
    return ledger


def formulas(book):
    return {f"{sheet.title}!{cell.coordinate}": cell.value
            for sheet in book for row in sheet for cell in row if cell.data_type == "f"}


def check_caches(path, expected, available=None):
    book = openpyxl.load_workbook(path, data_only=True)
    actual = {cell: book["Overview"][cell].value for cell in expected}
    require(actual == expected, f"Saved cache mismatch for {path}: {actual}")
    values = [book["Inventory"].cell(row, 6).value for row in range(8, 14)]
    if available is not None:
        require(values == available, f"Saved availability cache mismatch: {values}")
    for sheet in book:
        for row in sheet:
            for cell in row:
                require(cell.data_type != "e", f"Formula error in {sheet.title}!{cell.coordinate}")
    return {"overview": actual, "available": values}


def package_check(path):
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        prohibited = ("vbaproject", "externallinks/", "connections.xml", "embeddings/")
        require(not any(any(part in name.lower() for part in prohibited) for name in names),
                "Unexpected macro, external link, connection or embedding")
        for name in names:
            if name.endswith(".rels"):
                relationships = ET.fromstring(archive.read(name))
                require(not any(rel.get("TargetMode") == "External" for rel in relationships),
                        f"Unexpected external relationship in {name}")
        caches = {}
        for name in names:
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                for cell in root.findall(".//s:c", NS):
                    if cell.find("s:f", NS) is not None:
                        value = cell.find("s:v", NS)
                        require(value is not None and value.text is not None,
                                f"Missing saved formula cache in {name}:{cell.get('r')}")
                        caches[f"{name}:{cell.get('r')}"] = {"type": cell.get("t", "n"),
                                                               "value": value.text}
        require(len(caches) == 15, "Unexpected formula population")
        core = ET.fromstring(archive.read("docProps/core.xml"))
        for element in core:
            if element.tag.rsplit("}", 1)[-1] in {"creator", "lastModifiedBy"}:
                require(element.text in {None, "", "awesome-dot fictional example", "openpyxl"},
                        f"Unexpected document identity: {element.text!r}")
        return caches


def structure(book):
    inventory = book["Inventory"]
    return {"sheet_names": book.sheetnames,
            "sheet_states": {s.title: s.sheet_state for s in book},
            "names": {name: value.attr_text for name, value in book.defined_names.items()},
            "filter": inventory.auto_filter.ref,
            "freeze_panes": inventory.freeze_panes,
            "validation": [(str(v.sqref), v.type, v.operator, v.formula1, v.allow_blank)
                           for v in inventory.data_validations.dataValidation],
            "conditional_formats": [(str(key.sqref), [(r.type, r.operator, r.formula) for r in rules])
                                    for key, rules in inventory.conditional_formatting._cf_rules.items()],
            "tables": {s.title: list(s.tables) for s in book},
            "merged_ranges": {s.title: list(map(str, s.merged_cells.ranges)) for s in book},
            "hidden_rows": {s.title: [key for key, val in s.row_dimensions.items() if val.hidden] for s in book},
            "hidden_columns": {s.title: [key for key, val in s.column_dimensions.items() if val.hidden] for s in book},
            "print_areas": {s.title: str(s.print_area) for s in book},
            "show_gridlines": {s.title: s.sheet_view.showGridLines for s in book}}


def compare_workbooks(source, candidate, expected):
    before = openpyxl.load_workbook(source, data_only=False)
    after = openpyxl.load_workbook(candidate, data_only=False)
    require(formulas(before) == formulas(after), "Formula text changed")
    require(structure(before) == structure(after), "Workbook structure changed")
    actual_changes = []
    new_annotations = {item["cell"] for item in expected if item["rule"] == "REVIEW_EVIDENCE"}
    for sheet in before:
        other = after[sheet.title]
        for row in sheet:
            for cell in row:
                new = other[cell.coordinate]
                if cell.value != new.value or type(cell.value) is not type(new.value):
                    require(sheet.title == "Inventory", "Change outside authorized sheet")
                    actual_changes.append({"cell": cell.coordinate, "before": cell.value, "after": new.value})
                for attribute in ("font", "fill", "border", "alignment", "number_format", "protection"):
                    # Calc omits explicit font/alignment on empty note cells and
                    # resolves their inherited body style when text is inserted.
                    # New annotations must match the existing body cell in that row;
                    # all other cells must preserve their explicit style exactly.
                    comparison = (other[f"C{cell.row}"] if sheet.title == "Inventory"
                                  and cell.coordinate in new_annotations else cell)
                    require(copy(getattr(comparison, attribute)) == copy(getattr(new, attribute)),
                            f"Style changed at {sheet.title}!{cell.coordinate}: {attribute}")
    intended = [{key: item[key] for key in ("cell", "before", "after")} for item in expected]
    require(actual_changes == intended, f"Unexpected cell changes: {actual_changes}")
    inv = after["Inventory"]
    require([inv.cell(r, 1).value for r in range(8, 14)] == [f"export-A:{n}" for n in range(1, 7)],
            "Lineage/order changed")
    require([inv.cell(r, 2).value for r in range(8, 14)] == ["0007", "0012", "0025", "0031", "0044", "0044"],
            "Identifier values changed")
    require(all(inv.cell(r, 2).data_type == "s" and inv.cell(r, 2).number_format == "@"
                for r in range(8, 14)), "Identifier text type or format lost")
    require(inv["D9"].value == 0 and inv["D9"].data_type == "n", "Numeric zero lost")
    require(inv["D10"].value is None and inv["D11"].value is None, "Missing count became a value")
    return {"changed_cells": actual_changes, "formula_expressions": formulas(after),
            "structure": structure(after),
            "annotation_style": "Four newly populated Review reason cells match the existing body style; all other inspected cell styles are unchanged"}


def render(path, root, label, operations, soffice, pdftoppm):
    pdf = consumer_convert(path, root / (label + "-pdf"), "pdf",
                           root / (label + "-pdf-profile"), operations, soffice)
    preview_dir = root / (label + "-preview")
    preview_dir.mkdir()
    run_command([pdftoppm, "-scale-to", "1400", "-png", str(pdf),
                 str(preview_dir / "page")], operations)
    pages = sorted(preview_dir.glob("page-*.png"))
    require(len(pages) == 2, f"Expected one page per sheet, got {len(pages)}")
    return [str(p.relative_to(root)) for p in pages]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True, help="A directory that does not yet exist")
    parser.add_argument("--soffice", default="soffice")
    parser.add_argument("--pdftoppm", default="pdftoppm")
    args = parser.parse_args()
    output = args.output.absolute()
    require(not output.exists() and not output.is_symlink(), f"Refusing existing output directory: {output}")
    for executable in (args.soffice, args.pdftoppm):
        require(shutil.which(executable) is not None, f"Missing existing executable: {executable}")
    fixture = json.loads((HERE / "fixture.json").read_text(encoding="utf-8"))
    expected = json.loads((HERE / "expected.json").read_text(encoding="utf-8"))
    output.mkdir(parents=True)
    operations = []
    try:
        version = run_command([args.soffice, "--version"], operations)
        seed = output / "seed"
        seed.mkdir()
        make_source(seed / "inventory-source.xlsx", fixture)
        source = consumer_convert(seed / "inventory-source.xlsx", output / "source", "xlsx",
                                  output / "source-profile", operations, args.soffice)
        source_hash = digest(source)
        source_caches = check_caches(source, expected["source_caches"])
        source_preview = render(source, output, "source", operations, args.soffice, args.pdftoppm)
        edited = output / "edited"
        edited.mkdir()
        ledger = transform(source, edited / "inventory-cleaned.xlsx", fixture)
        require(ledger == expected["changes"], "Transformation ledger differs from independent expected decisions")
        candidate = consumer_convert(edited / "inventory-cleaned.xlsx", output / "candidate", "xlsx",
                                     output / "candidate-profile", operations, args.soffice)
        candidate_caches = check_caches(candidate, expected["candidate_caches"], expected["candidate_available"])
        saved_checks = compare_workbooks(source, candidate, expected["changes"])
        xml_caches = package_check(candidate)
        package_check(source)
        # Independent control: literal whole-unit arithmetic, separate from both the
        # transformation and spreadsheet formulas. Missing rows stay out of both sides.
        require(4 + 0 + 3 + 5 == expected["reconciliation"]["known_ready_units"], "Ready-unit control failed")
        require(1 + 0 + 1 + 1 == expected["reconciliation"]["reserved_units_on_rows_with_known_ready"], "Reserved-unit control failed")
        require((4-1) + (0-0) + (3-1) + (5-1) == candidate_caches["overview"]["B4"], "Availability control failed")
        # A real input change is recalculated and saved by the consumer. The delivered
        # candidate is never changed. Zero ready minus one reserved is a -1 shortage.
        probe_input = output / "probe-input"
        probe_input.mkdir()
        probe_book = openpyxl.load_workbook(candidate, data_only=False)
        probe_book["Inventory"]["D10"] = 0
        probe_book.save(probe_input / "zero-probe.xlsx")
        probe = consumer_convert(probe_input / "zero-probe.xlsx", output / "probe-result", "xlsx",
                                 output / "probe-profile", operations, args.soffice)
        probe_caches = check_caches(probe, expected["probe_caches"], expected["probe_available"])
        require(formulas(openpyxl.load_workbook(candidate)) == formulas(openpyxl.load_workbook(probe)),
                "Probe changed formulas")
        package_check(probe)
        candidate_preview = render(candidate, output, "candidate", operations, args.soffice, args.pdftoppm)
        require(digest(source) == source_hash, "Source changed after capture")
        artifacts = {str(path.relative_to(output)): digest(path) for path in [source, candidate]}
        for path in output.rglob("*"):
            if path.suffix.lower() in {".xlsx", ".png", ".pdf"}:
                require(path.stat().st_size <= 2 * 1024 * 1024, f"Binary artifact exceeds 2 MiB: {path}")
        report = {"result": "pass", "checked_at_utc": datetime.now(timezone.utc).isoformat(),
                  "engine": version, "python": sys.version.split()[0],
                  "openpyxl": openpyxl.__version__, "xlsxwriter": xlsxwriter.__version__,
                  "source_preserved": True, "artifacts": artifacts,
                  "source_caches": source_caches, "candidate_caches": candidate_caches,
                  "probe": {"change": "Inventory!D10: blank to numeric 0 in a disposable copy",
                            "saved_caches": probe_caches},
                  "saved_file_checks": saved_checks, "raw_xml_formula_caches": xml_caches,
                  "reconciliation": expected["reconciliation"],
                  "previews": {"source": source_preview, "candidate": candidate_preview},
                  "operations": operations,
                  "limits": ["Fictional six-row local XLSX fixture only; no independent business control",
                             "LibreOffice version is the engine above; Microsoft Excel and live Google Sheets were not tested",
                             "No macros, external links, pivots, charts, native tables, protection or data connections are present or tested",
                             "Verification covers saved formulas, caches and listed structure; it is not a general workbook-fidelity guarantee",
                             "PNG previews need visual review; automated success is not visual approval"]}
        dump(output / "report.json", report)
        print(json.dumps({"result": "pass", "report": str(output / "report.json"),
                          "artifacts": artifacts, "candidate_previews": candidate_preview}, indent=2))
    except Exception as error:
        dump(output / "failure.json", {"error": str(error), "operations": operations})
        raise


if __name__ == "__main__":
    main()
